"""Génère les Word courants depuis les Markdown, sans texte dupliqué.

Usage : python docs/build_word.py
"""
from __future__ import annotations

from pathlib import Path
import json
import re
from urllib.parse import quote, urlsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

HERE = Path(__file__).resolve().parent
OUT = HERE / "word"
DOCUMENTS = json.loads((HERE / "documents.json").read_text(encoding="utf-8"))
CURRENT_SOURCE = HERE
# Tailles de lecture pour les figures des tutoriels (largeur, hauteur maxi).
FIGURE_SIZES = {
    "08-sessions": (17.1, 10.8),
    "05k-point-ponctuel-detail": (8.0, 12.0),
    "05m-trajectoire-in-detail": (6.8, 12.0),
    "05j-broutage-et-bouchees-detail": (8.8, 10.8),
}


def inline(p, text):
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    for token in re.split(pattern, text):
        if token.startswith("**") and token.endswith("**"):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            p.add_run(token[1:-1]).font.name = "Consolas"
        elif (match := re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)):
            link = OxmlElement("w:hyperlink")
            target = match[2]
            parts = urlsplit(target)
            if not parts.scheme and not parts.netloc:
                local = (CURRENT_SOURCE.parent / parts.path).resolve()
                relative = local.relative_to(HERE.parent).as_posix()
                target = "https://github.com/sas-ldi/AquaMeasure/blob/main/" + quote(relative)
                if parts.fragment:
                    target += "#" + parts.fragment
            link.set(qn("r:id"), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "155B86")
            props.append(color); run.append(props)
            content = OxmlElement("w:t"); content.text = match[1]
            run.append(content); link.append(run); p._p.append(link)
        else:
            p.add_run(token)


def new_document(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.header_distance = sec.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string("4F5962")
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.space_after = Pt(4)
    header = sec.header.paragraphs[0]
    header.text = "AquaMeasure"
    header.style = doc.styles["Caption"]
    for run in header.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Version du 8 septembre 2026   ·   ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_paragraph(title, "Title")
    doc.core_properties.title = title
    doc.core_properties.author = "AquaMeasure"
    return doc


def math_run(text):
    node = OxmlElement("m:r")
    content = OxmlElement("m:t")
    content.text = text
    node.append(content)
    return node


def math_container(tag, *children):
    node = OxmlElement("m:" + tag)
    for child in children:
        node.append(math_run(child) if isinstance(child, str) else child)
    return node


def sub(base, index):
    return math_container("sSub", math_container("e", base), math_container("sub", index))


def sup(base, power):
    return math_container("sSup", math_container("e", base), math_container("sup", power))


def fraction(numerator, denominator):
    return math_container("f", math_container("num", *numerator), math_container("den", *denominator))


def equation(doc, *children):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(8)
    p._p.append(math_container("oMathPara", math_container("oMath", *children)))


def render_equation(doc, code):
    if code.startswith("s [u, v, 1]"):
        equation(doc, "s ", sup("[u, v, 1]", "T"), " = K [R | t] ", sup("[X, Y, Z, 1]", "T"))
    elif code.startswith("d = u_gauche"):
        equation(doc, "d = ", sub("u", "gauche"), " − ", sub("u", "droite"))
        equation(doc, "Z = ", fraction(["f × B"], ["d"]))
    elif code.startswith("L = √"):
        terms = []
        for axis in "XYZ":
            if terms:
                terms.append(" + ")
            term = math_container("e", "(", sub(axis, "B"), " − ", sub(axis, "A"), ")")
            terms.append(math_container("sSup", term, math_container("sup", "2")))
        props = math_container("radPr")
        hide = OxmlElement("m:degHide"); hide.set(qn("m:val"), "1"); props.append(hide)
        radical = math_container("rad", props, math_container("deg"), math_container("e", *terms))
        equation(doc, "L = ", radical)
    elif code.startswith("σZ ≈"):
        equation(doc, sub("σ", "Z"), " ≈ ", fraction([sup("Z", "2"), " × ", sub("σ", "d")], ["f × B"]))
    else:
        return False
    return True


def render(doc, source, *, skip_title=True):
    global CURRENT_SOURCE
    CURRENT_SOURCE = source
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1; continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            if render_equation(doc, "\n".join(code)):
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = len(code) < 15
            run = p.add_run("\n".join(code)); run.font.name = "Consolas"; run.font.size = Pt(9)
            i += 1; continue
        if (match := re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)):
            image = source.parent / match[2]
            if not image.is_file():
                raise FileNotFoundError(image)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            with Image.open(image) as bitmap:
                width, height = bitmap.size
            # Les gros plans de cartes gardent une taille lisible, sans
            # occuper une page entière comme une capture panoramique.
            if "-detail" in image.stem:
                target_width = 8.8 if width < 700 else (10.6 if width < 1200 else 17.1)
                max_height = 12.0 if width < 1200 else 18.0
            else:
                target_width, max_height = 17.1, 12.0
            target_width, max_height = FIGURE_SIZES.get(image.stem, (target_width, max_height))
            width_cm = min(target_width, max_height * width / height)
            shape = p.add_run().add_picture(str(image), width=Cm(width_cm))
            shape._inline.docPr.set("descr", match[1])
            cap = doc.add_paragraph(match[1], "Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                    rows.append(row)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            if len(rows[0]) == 2:
                table.autofit = False
                table.columns[0].width = Cm(5.6)
                table.columns[1].width = Cm(11.6)
            props = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement("w:tblCellMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                if len(cells) == 2:
                    cells[0].width = Cm(5.6)
                    cells[1].width = Cm(11.6)
                for cell, content in zip(cells, row):
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    if len(rows) <= 8 and index < len(rows) - 1:
                        cell.paragraphs[0].paragraph_format.keep_with_next = True
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EDEFF2" if index == 0 else ("F8F9FA" if index % 2 == 0 else "FFFFFF"))
                    cell._tc.get_or_add_tcPr().append(shade)
                    inline(cell.paragraphs[0], content)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9.5)
                        if index == 0: run.bold = True
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if index == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if (match := re.match(r"^(#{1,4}) (.+)", line)):
            if len(match[1]) == 1 and skip_title:
                i += 1; continue
            doc.add_heading(match[2], max(1, len(match[1]) - 1))
            i += 1; continue
        if (match := re.match(r"^(\d+\.|[-*]) (.+)", line)):
            # Garder le numéro explicite évite que Word poursuive la liste
            # d'une section précédente.
            p = doc.add_paragraph(style="List Bullet" if match[1] in "-*" else "Normal")
            if match[1] not in "-*":
                p.paragraph_format.left_indent = Cm(0.35)
                p.paragraph_format.first_line_indent = Cm(-0.35)
                inline(p, match[1] + " " + match[2])
            else:
                inline(p, match[2])
            # Garder les étapes d'un geste ensemble quand elles tiennent
            # sur une page ; un tutoriel ne commence pas au point 3.
            next_index = i + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            if next_index < len(lines) and re.match(r"^\d+\. ", lines[next_index]):
                p.paragraph_format.keep_with_next = True
            i += 1; continue
        parts = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|!\[|\||```|\d+\. |[-*] )", lines[i]):
            parts.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        inline(p, " ".join(parts))
        next_index = i
        while next_index < len(lines) and not lines[next_index].strip():
            next_index += 1
        if next_index < len(lines) and re.match(r"^(\d+\. |\||```)", lines[next_index]):
            p.paragraph_format.keep_with_next = True


def main():
    OUT.mkdir(exist_ok=True)
    for entry in DOCUMENTS:
        source = HERE / entry["source"]
        output = OUT / entry["category"] / (entry["stem"] + ".docx")
        output.parent.mkdir(parents=True, exist_ok=True)
        title = source.read_text(encoding="utf-8-sig").splitlines()[0].lstrip("# ")
        doc = new_document(title)
        render(doc, source)
        doc.save(output)
        print(output.relative_to(HERE), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
