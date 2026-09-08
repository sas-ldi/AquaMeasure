"""Génère les Word courants depuis les Markdown, sans texte dupliqué.

Usage : python aquameasure-pyside/docs/build_word.py
"""
from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

HERE = Path(__file__).resolve().parent
OUT = HERE / "word"
DOCUMENTS = [
    ("MANUEL_UTILISATEUR.md", "AquaMeasure_Manuel_Utilisateur.docx"),
    ("DETECTEURS.md", "AquaMeasure_Guide_Extension_Modeles_Detection.docx"),
]
# Tailles de lecture pour les figures des tutoriels (largeur, hauteur maxi).
FIGURE_SIZES = {
    "08-sessions": (17.1, 10.8),
    "05k-point-ponctuel-detail": (8.0, 12.0),
    "05m-trajectoire-in-detail": (6.8, 12.0),
    "05j-broutage-et-bouchees-detail": (8.8, 10.8),
}


def inline(p, text):
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    for token in re.split(pattern, text):
        if token.startswith("**") and token.endswith("**"):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            p.add_run(token[1:-1]).font.name = "Consolas"
        elif (match := re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)):
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), p.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "155B86")
            props.append(color); run.append(props)
            content = OxmlElement("w:t"); content.text = match[1]
            run.append(content); link.append(run); p._p.append(link)
        else:
            p.add_run(token)


def new_document(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.header_distance = sec.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string("4F5962")
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.space_after = Pt(4)
    header = sec.header.paragraphs[0]
    header.text = "AquaMeasure"
    header.style = doc.styles["Caption"]
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Version du 8 septembre 2026   ·   ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_paragraph(title, "Title")
    doc.core_properties.title = title
    doc.core_properties.author = "AquaMeasure"
    return doc


def render(doc, source, *, skip_title=True):
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1; continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = len(code) < 15
            run = p.add_run("\n".join(code)); run.font.name = "Consolas"; run.font.size = Pt(9)
            i += 1; continue
        if (match := re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)):
            image = source.parent / match[2]
            if not image.is_file():
                raise FileNotFoundError(image)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            with Image.open(image) as bitmap:
                width, height = bitmap.size
            # Les gros plans de cartes gardent une taille lisible, sans
            # occuper une page entière comme une capture panoramique.
            if "-detail" in image.stem:
                target_width = 8.8 if width < 700 else (10.6 if width < 1200 else 17.1)
                max_height = 12.0 if width < 1200 else 18.0
            else:
                target_width, max_height = 17.1, 12.0
            target_width, max_height = FIGURE_SIZES.get(image.stem, (target_width, max_height))
            width_cm = min(target_width, max_height * width / height)
            shape = p.add_run().add_picture(str(image), width=Cm(width_cm))
            shape._inline.docPr.set("descr", match[1])
            cap = doc.add_paragraph(match[1], "Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                    rows.append(row)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            props = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement("w:tblCellMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, content in zip(cells, row):
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    if len(rows) <= 8 and index < len(rows) - 1:
                        cell.paragraphs[0].paragraph_format.keep_with_next = True
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EDEFF2" if index == 0 else ("F8F9FA" if index % 2 == 0 else "FFFFFF"))
                    cell._tc.get_or_add_tcPr().append(shade)
                    inline(cell.paragraphs[0], content)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9.5)
                        if index == 0: run.bold = True
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if index == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if (match := re.match(r"^(#{1,4}) (.+)", line)):
            if len(match[1]) == 1 and skip_title:
                i += 1; continue
            doc.add_heading(match[2], max(1, len(match[1]) - 1))
            i += 1; continue
        if (match := re.match(r"^(\d+\.|[-*]) (.+)", line)):
            # Garder le numéro explicite évite que Word poursuive la liste
            # d'une section précédente.
            p = doc.add_paragraph(style="List Bullet" if match[1] in "-*" else "Normal")
            if match[1] not in "-*":
                p.paragraph_format.left_indent = Cm(0.35)
                p.paragraph_format.first_line_indent = Cm(-0.35)
                inline(p, match[1] + " " + match[2])
            else:
                inline(p, match[2])
            # Garder les étapes d'un geste ensemble quand elles tiennent
            # sur une page ; un tutoriel ne commence pas au point 3.
            next_index = i + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            if next_index < len(lines) and re.match(r"^\d+\. ", lines[next_index]):
                p.paragraph_format.keep_with_next = True
            i += 1; continue
        parts = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|!\[|\||```|\d+\. |[-*] )", lines[i]):
            parts.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        inline(p, " ".join(parts))
        next_index = i
        while next_index < len(lines) and not lines[next_index].strip():
            next_index += 1
        if next_index < len(lines) and re.match(r"^(\d+\. |\|)", lines[next_index]):
            p.paragraph_format.keep_with_next = True


def main():
    OUT.mkdir(exist_ok=True)
    for source_name, output in DOCUMENTS:
        source = HERE / source_name
        title = source.read_text(encoding="utf-8-sig").splitlines()[0].lstrip("# ")
        doc = new_document(title)
        render(doc, source)
        doc.save(OUT / output)
        print(output, flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
